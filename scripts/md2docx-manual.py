# -*- coding: utf-8 -*-
# 用户手册 docx 生成：从 docs/使用说明.md 生成运营商公文风格 Word 版
# 格式：标题黑体18pt居中 / 章节黑体14pt / 正文仿宋12pt首行缩进2字符 / 1.5倍行距 / 表头黑体加粗
import re, sys
from docx import Document
from docx.shared import Pt, Emu
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITLE_FONT, HEADER_FONT, BODY_FONT = '黑体', '黑体', '仿宋'
TITLE_SIZE, HEADER_SIZE, BODY_SIZE = Pt(18), Pt(14), Pt(12)
LINE_SPACING, FIRST_LINE_INDENT = 1.5, Pt(24)

def set_run_font(run, font_name, font_size, bold=False):
    run.font.size = font_size
    run.font.bold = bold
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {}); rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    rfonts.set(qn('w:ascii'), font_name)
    rfonts.set(qn('w:hAnsi'), font_name)

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LINE_SPACING
    p.paragraph_format.first_line_indent = FIRST_LINE_INDENT
    r = p.add_run(text)
    set_run_font(r, BODY_FONT, BODY_SIZE)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LINE_SPACING
    r = p.add_run(text)
    set_run_font(r, TITLE_FONT, TITLE_SIZE, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = LINE_SPACING
    r = p.add_run(text)
    set_run_font(r, HEADER_FONT, HEADER_SIZE, bold=True)
    return p

def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(cell_text)
            set_run_font(r, '黑体' if ri == 0 else BODY_FONT, Pt(12), bold=(ri == 0))
    return t

# —— 解析 markdown → docx ——
src = open(sys.argv[1], encoding='utf-8').read()
doc = Document()
sec = doc.sections[0]
sec.page_width = Emu(7772400); sec.page_height = Emu(10058400)
sec.left_margin = Emu(1141095); sec.right_margin = Emu(1141095)
sec.top_margin = Emu(914400); sec.bottom_margin = Emu(791845)

lines = src.split('\n')
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1; continue
    if line.startswith('# '):
        add_h1(doc, line[2:].strip()); i += 1; continue
    if line.startswith('## '):
        add_h2(doc, line[3:].strip()); i += 1; continue
    if line.startswith('### '):
        # 三级标题：仿宋加粗，首行缩进保留
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = LINE_SPACING
        p.paragraph_format.first_line_indent = FIRST_LINE_INDENT
        r = p.add_run(line[4:].strip())
        set_run_font(r, BODY_FONT, BODY_SIZE, bold=True)
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            raw = lines[i].strip()
            if re.match(r'^\|[\s\-:|]+\|$', raw):
                i += 1; continue
            cells = [c.strip() for c in raw.strip('|').split('|')]
            rows.append(cells); i += 1
        if rows: add_table(doc, rows)
        continue
    if line.startswith('```'):
        buf = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            buf.append(lines[i]); i += 1
        i += 1
        for cl in buf:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(cl if cl else ' ')
            set_run_font(r, 'Consolas', Pt(9))
        continue
    if line.startswith('>'):
        # 引用块：连续 > 行合并为一段（去掉 > 前缀与空行占位 >）
        buf = []
        while i < len(lines) and lines[i].startswith('>'):
            seg = lines[i].strip()
            if seg.startswith('>'): seg = seg[1:].lstrip()
            if seg: buf.append(seg)
            i += 1
        add_body(doc, ''.join(buf).replace('`', ''))
        continue
    # 列表项：- / 数字. 前缀去掉，保留正文
    text = re.sub(r'^\s*[-*]\s+', '', line)
    text = re.sub(r'^\s*\d+\.\s+', '', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = text.replace('`', '')
    add_body(doc, text); i += 1

doc.save(sys.argv[2])
print('saved:', sys.argv[2])
